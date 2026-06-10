import asyncio
import logging
import os
import tempfile
from typing import Sequence, Iterable, Any
from datetime import datetime, date

from aiogram import Router, F
from aiogram.filters import StateFilter
from aiogram.fsm.context import FSMContext
from aiogram.types import Message
from aiogram.types.input_file import FSInputFile

from sqlalchemy import select
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

import app.admin.admin_keyboard as kb
from app.users.keyboard import get_main_keyboard
from app.vpn.provisioning import PROTOCOL_WIREGUARD, PROTOCOL_XUI, get_vpn_client_count
from config import ADMIN_ID
from app.database.models import async_session, TestPeriod, User, Subscribers, Payments, Server

admin_router = Router()
logger = logging.getLogger(__name__)

DATE_FMT = "%Y-%m-%d"


def is_admin(user_id: int) -> bool:
    return user_id == int(ADMIN_ID)


def _admin_actor(user) -> str:
    username = user.username or "-"
    return f"{user.id} (@{username})"


def autosize_columns(ws):
    for col_idx, col in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in col:
            if cell.value is None:
                continue
            max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 50)


def safe_parse_date(d: str) -> date | None:
    try:
        return datetime.strptime(str(d), DATE_FMT).date()
    except Exception:
        return None


async def send_excel(
    message: Message,
    filename: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
    sheet_name: str = "data",
):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    ws.append(list(headers))
    for r in rows:
        ws.append(list(r))

    autosize_columns(ws)

    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        tmp_path = tmp.name

    try:
        wb.save(tmp_path)
        doc = FSInputFile(tmp_path, filename=filename)
        await message.answer_document(doc)
    finally:
        try:
            os.remove(tmp_path)
        except Exception:
            logger.exception("Не удалось удалить временный файл отчета: path=%s", tmp_path)


@admin_router.message(F.text == "Админ")
async def admin_panel_button(message: Message, state: FSMContext):
    if is_admin(message.from_user.id):
        await state.clear()
        await message.answer("Вы вошли в админ-панель", reply_markup=kb.admin_panel)
    else:
        await message.answer("У вас нет доступа")


@admin_router.message(F.text == "Назад (Админ)")
async def help_main_button(message: Message, state: FSMContext):
    if is_admin(message.from_user.id):
        await state.clear()
        await message.answer("Вы вернулись в главное меню", reply_markup=get_main_keyboard(message.from_user.id))
    else:
        await message.answer("У вас нет доступа")


@admin_router.message(F.text == "Статистика")
async def stats_menu(message: Message):
    if is_admin(message.from_user.id):
        await message.answer("Выберите параметр, который вас интересует", reply_markup=kb.stat_kb)
    else:
        await message.answer("У вас нет доступа")


@admin_router.message(F.text == "Назад Админ")
async def back_admin(message: Message, state: FSMContext):
    if is_admin(message.from_user.id):
        await state.clear()
        await message.answer("Вы вернулись в главное меню", reply_markup=kb.admin_panel)
    else:
        await message.answer("У вас нет доступа")


@admin_router.message(
    StateFilter("*"),
    F.from_user.id == int(ADMIN_ID),
    F.text.in_({"❌ Отмена", "Отмена", "отмена"}),
)
async def cancel_admin_action(message: Message, state: FSMContext):
    current_state = await state.get_state()
    if current_state is None:
        await message.answer("Нет активного действия для отмены.", reply_markup=kb.admin_panel)
        return

    await state.clear()
    await message.answer("Действие отменено.", reply_markup=kb.admin_panel)


@admin_router.message(F.text == "Клиенты на сервере")
async def clients_on_servers_wg(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("У вас нет доступа")
        return

    async with async_session() as session:
        res = await session.execute(
            select(Server)
            .where(Server.is_active == True)
            .order_by(Server.region, Server.region_id)
        )
        servers = res.scalars().all()

    total_clients = 0
    lines = ["📊 <b>Клиенты VPN по сервисам</b>\n", "<b>WireGuard:</b>"]

    if not servers:
        lines.append("• нет активных серверов")
    else:
        for s in servers:
            try:
                count = await get_vpn_client_count(PROTOCOL_WIREGUARD, s)
                total_clients += int(count)
                lines.append(f"• <b>{s.region} №{s.region_id}</b>: <b>{count}</b>")
            except Exception:
                logger.exception(
                    "Ошибка запроса клиентов WG: admin=%s region=%s region_id=%s host=%s",
                    _admin_actor(message.from_user),
                    s.region,
                    s.region_id,
                    s.host_ip,
                )
                lines.append(f"• <b>{s.region} №{s.region_id}</b>: ⚠️ ошибка")

    try:
        xui_count = await get_vpn_client_count(PROTOCOL_XUI)
        total_clients += int(xui_count)
        lines.append(f"\n<b>3xUI:</b>\n• <b>Панель</b>: <b>{xui_count}</b>")
    except Exception:
        logger.exception("Ошибка запроса клиентов 3xUI: admin=%s", _admin_actor(message.from_user))
        lines.append("\n<b>3xUI:</b>\n• <b>Панель</b>: ⚠️ ошибка")

    lines.append(f"\n<b>Итого клиентов:</b> {total_clients}")

    logger.info(
        "Админ %s получил статистику клиентов VPN: серверов=%s всего_клиентов=%s",
        _admin_actor(message.from_user),
        len(servers),
        total_clients,
    )
    await message.answer("\n".join(lines), parse_mode="HTML", reply_markup=kb.stat_kb)

# ====== Users -> Excel + summary ======
@admin_router.message(F.text == "Пользователи в боте")
async def users_excel(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("У вас нет доступа")
        return

    async with async_session() as session:
        res = await session.execute(
            select(User.id, User.tg_id, User.username, User.first_name, User.date).order_by(User.id)
        )
        rows = res.all()

    total = len(rows)
    await message.answer(
        f"📊 <b>Users</b>\n"
        f"Всего пользователей: <b>{total}</b>\n"
        f"Файл: <b>users.xlsx</b>",
        parse_mode="HTML",
    )

    await send_excel(
        message=message,
        filename="users.xlsx",
        sheet_name="users",
        headers=["id", "tg_id", "username", "first_name", "date"],
        rows=rows,
    )
    logger.info(
        "Админ %s выгрузил users.xlsx: rows=%s",
        _admin_actor(message.from_user),
        total,
    )


# ====== Subscribers -> Excel + summary ======
@admin_router.message(F.text == "Пользователи с подпиской")
async def subscribers_excel(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("У вас нет доступа")
        return

    async with async_session() as session:
        res = await session.execute(
            select(
                Subscribers.id,
                Subscribers.tg_id,
                Subscribers.username,
                Subscribers.file_name,
                Subscribers.subscription,
                Subscribers.expiry_date,
                Subscribers.server_region,
                Subscribers.server_region_id,
                Subscribers.protocol,
                Subscribers.xui_sub_id,
                Subscribers.notif_oneday,
                Subscribers.note,
            ).order_by(Subscribers.id)
        )
        rows = res.all()

    today = datetime.now().date()
    total = len(rows)

    active = 0
    expired = 0
    bad_dates = 0

    by_region: dict[str, int] = {}
    for r in rows:
        # r: (id, tg_id, username, file_name, subscription, expiry_date, server_region, server_region_id, protocol, xui_sub_id, notif_oneday, note)
        exp = safe_parse_date(r[5])
        if exp is None:
            bad_dates += 1
        else:
            if exp >= today:
                active += 1
            else:
                expired += 1

        region_key = f"{r[8]} | {r[6]} #{r[7]}"
        by_region[region_key] = by_region.get(region_key, 0) + 1

    top_regions = sorted(by_region.items(), key=lambda x: x[1], reverse=True)[:5]
    top_regions_text = "\n".join([f"• {k}: <b>{v}</b>" for k, v in top_regions]) or "—"

    await message.answer(
        f"📊 <b>Subscribers</b>\n"
        f"Всего подписок: <b>{total}</b>\n"
        f"Активных (expiry_date ≥ сегодня): <b>{active}</b>\n"
        f"Просроченных: <b>{expired}</b>\n"
        f"Некорректных дат: <b>{bad_dates}</b>\n\n"
        f"<b>Топ серверов (подписок):</b>\n{top_regions_text}\n\n"
        f"Файл: <b>subscribers.xlsx</b>",
        parse_mode="HTML",
    )

    await send_excel(
        message=message,
        filename="subscribers.xlsx",
        sheet_name="subscribers",
        headers=[
            "id", "tg_id", "username", "file_name", "subscription",
            "expiry_date", "server_region", "server_region_id",
            "protocol", "xui_sub_id", "notif_oneday", "note"
        ],
        rows=rows,
    )
    logger.info(
        "Админ %s выгрузил subscribers.xlsx: total=%s active=%s expired=%s bad_dates=%s",
        _admin_actor(message.from_user),
        total,
        active,
        expired,
        bad_dates,
    )


# ====== TestPeriod -> Excel + summary ======
@admin_router.message(F.text == "Пользователи с пробным периодом")
async def testperiod_excel(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("У вас нет доступа")
        return

    async with async_session() as session:
        res = await session.execute(
            select(
                TestPeriod.id,
                TestPeriod.tg_id,
                TestPeriod.username,
                TestPeriod.file_name,
                TestPeriod.subscription,
                TestPeriod.expiry_date,
                TestPeriod.protocol,
                TestPeriod.xui_sub_id,
                TestPeriod.notif_oneday,
            ).order_by(TestPeriod.id)
        )
        rows = res.all()

    today = datetime.now().date()
    total = len(rows)

    active = 0
    expired = 0
    bad_dates = 0

    for r in rows:
        exp = safe_parse_date(r[5])
        if exp is None:
            bad_dates += 1
        else:
            if exp >= today:
                active += 1
            else:
                expired += 1

    await message.answer(
        f"📊 <b>Test period</b>\n"
        f"Всего записей: <b>{total}</b>\n"
        f"Активных: <b>{active}</b>\n"
        f"Просроченных: <b>{expired}</b>\n"
        f"Некорректных дат: <b>{bad_dates}</b>\n\n"
        f"Файл: <b>test_period.xlsx</b>",
        parse_mode="HTML",
    )

    await send_excel(
        message=message,
        filename="test_period.xlsx",
        sheet_name="test_period",
        headers=["id", "tg_id", "username", "file_name", "subscription", "expiry_date", "protocol", "xui_sub_id", "notif_oneday"],
        rows=rows,
    )
    logger.info(
        "Админ %s выгрузил test_period.xlsx: total=%s active=%s expired=%s bad_dates=%s",
        _admin_actor(message.from_user),
        total,
        active,
        expired,
        bad_dates,
    )


# ====== Payments -> Excel + summary ======
@admin_router.message(F.text == "Платежи")
async def payments_excel(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("У вас нет доступа")
        return

    async with async_session() as session:
        res = await session.execute(
            select(
                Payments.id,
                Payments.tg_id,
                Payments.username,
                Payments.price,
                Payments.date,
                Payments.tarific_plan,
                Payments.provider_payment_charge_id,
            ).order_by(Payments.id)
        )
        rows = res.all()

    total = len(rows)
    total_sum = 0.0
    count_price = 0

    by_plan: dict[str, int] = {}
    for r in rows:
        price = r[3]
        try:
            p = float(price)
            total_sum += p
            count_price += 1
        except Exception:
            pass

        plan = str(r[5])
        by_plan[plan] = by_plan.get(plan, 0) + 1

    avg = (total_sum / count_price) if count_price else 0.0

    top_plans = sorted(by_plan.items(), key=lambda x: x[1], reverse=True)[:5]
    top_plans_text = "\n".join([f"• {k}: <b>{v}</b>" for k, v in top_plans]) or "—"

    await message.answer(
        f"📊 <b>Payments</b>\n"
        f"Всего платежей: <b>{total}</b>\n"
        f"Сумма (price): <b>{total_sum:.2f}</b>\n"
        f"Средний чек: <b>{avg:.2f}</b>\n\n"
        f"<b>Топ тарифов (по кол-ву платежей):</b>\n{top_plans_text}\n\n"
        f"Файл: <b>payments.xlsx</b>",
        parse_mode="HTML",
    )

    await send_excel(
        message=message,
        filename="payments.xlsx",
        sheet_name="payments",
        headers=[
            "id", "tg_id", "username", "price", "date",
            "tarific_plan", "provider_payment_charge_id"
        ],
        rows=rows,
    )
    logger.info(
        "Админ %s выгрузил payments.xlsx: total=%s sum=%.2f avg=%.2f",
        _admin_actor(message.from_user),
        total,
        total_sum,
        avg,
    )
